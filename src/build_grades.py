"""Encrypt the department grades workbook into the site (one record per student, locked with that
student's personal number: PBKDF2-SHA256 600k -> AES-GCM). The workbook itself never goes into the repo.

Usage (from the repo root):
  python3 src/build_grades.py <workbook.xlsx> "<update label>" [--inject]

- Reads only the class sheets present in the workbook ("الأول 1 تربية" + "الأول 1 تلاوة" …).
  Classes that are not in the workbook keep their current records from src/gr.json untouched.
- Names must match src/st.json row by row; any mismatch stops the build.
- Empty cell = not recorded yet (shown on the site as «يظهر لاحقًا»).
- A 0 in حفظ/تلاوة is kept as 0; the site shows it as «لم يقم بالتسميع» with a note to see the teacher.
- --inject writes the result into index.html (var GR=…) as well as src/gr.json.
"""
import sys, json, os, base64, hashlib, re
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

ITER = 600000
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GN = {'السادس': '6', 'الأول': '7', 'الثاني': '8', 'الثالث': '9'}
ST = json.load(open(os.path.join(ROOT, 'src', 'st.json'), encoding='utf-8'))
OLD = json.load(open(os.path.join(ROOT, 'src', 'gr.json'), encoding='utf-8'))


def num(v):
    if v is None or v == '':
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def b64(b):
    return base64.b64encode(b).decode()


wb = openpyxl.load_workbook(sys.argv[1], data_only=False)
label = sys.argv[2] if len(sys.argv) > 2 else ''
out = dict(OLD['d'])
done = []
for key, names in ST.items():
    g, s = key.split('-')
    lab = [k for k, v in GN.items() if v == g][0] + ' ' + s
    if lab + ' تربية' not in wb.sheetnames or lab + ' تلاوة' not in wb.sheetnames:
        continue
    wi, wt = wb[lab + ' تربية'], wb[lab + ' تلاوة']
    recs = []
    for i, n in enumerate(names):
        ri, rt = 5 + i, 4 + i
        if norm(wi.cell(ri, 3).value) != norm(n) or norm(wt.cell(rt, 3).value) != norm(n):
            raise SystemExit(f'name mismatch {lab} row {i + 1}')
        v = lambda c: num(wi.cell(ri, c).value)
        w = lambda c: num(wt.cell(rt, c).value)
        rec = {'n': n,
               'i': {'mul': [v(4), v(5), v(6), v(7)],      # D..G  وقفة 1/2 سلوك وأداء
                     'port': [v(9), v(10), v(11)],         # I..K  ملف الأعمال
                     'task': [v(13), v(14), v(15)],        # M..O  المهمة
                     'exam': [v(17), v(18)]},              # Q..R  الاختبار 1/2
               't': {'hifz': [w(4), w(5)],                 # D..E  الحفظ 1/2
                     'tila': [w(8), w(9), w(10), w(11)]}}  # H..K  التلاوة رصد 1/2
        code = re.sub(r'\D', '', str(wi.cell(ri, 2).value or ''))
        if len(code) != 9:
            recs.append(None)
            continue
        salt, iv = os.urandom(16), os.urandom(12)
        k = hashlib.pbkdf2_hmac('sha256', code.encode(), salt, ITER, 32)
        ct = AESGCM(k).encrypt(iv, json.dumps(rec, ensure_ascii=False).encode(), None)
        recs.append([b64(salt), b64(iv), b64(ct)])
    out[key] = recs
    done.append(f'{key}: {sum(1 for r in recs if r)}/{len(recs)}')

GR = {'iter': ITER, 'd': out, 'updated': label}
js = json.dumps(GR, ensure_ascii=False, separators=(',', ':'))
open(os.path.join(ROOT, 'src', 'gr.json'), 'w', encoding='utf-8').write(js)
if '--inject' in sys.argv:
    for p in ('index.html',):
        f = os.path.join(ROOT, p)
        h = open(f, encoding='utf-8').read()
        a = h.index('var GR=')
        b = h.index(';\n', a)
        open(f, 'w', encoding='utf-8').write(h[:a] + 'var GR=' + js + h[b:])
print('updated classes →', ', '.join(done))
