import openpyxl, json, copy, secrets, sys
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
TPL='/root/.claude/uploads/3135d92f-e321-5b42-b747-2692c15712ae/c8f51866-____________________-___________.xlsx'
ST=json.load(open('/home/claude/work/st.json',encoding='utf-8'))
GN={'7':'الأول','8':'الثاني','9':'الثالث'}
wb=openpyxl.load_workbook(TPL)
src_i=wb['التربية الاسلامية']; src_t=wb['التلاوة']
def clone(src,title):
    ws=wb.copy_worksheet(src); ws.title=title
    ws.sheet_view.rightToLeft=True
    for dv in src.data_validations.dataValidation:
        n=DataValidation(type=dv.type,operator=dv.operator,formula1=dv.formula1,formula2=dv.formula2,allow_blank=dv.allow_blank,showErrorMessage=True,error=dv.error)
        n.sqref=dv.sqref; ws.add_data_validation(n)
    ws.protection=copy.copy(src.protection)
    return ws
for key in sorted(ST):
    g,s=key.split('-'); names=ST[key]; lab=f'{GN[g]} {s}'
    wi=clone(src_i,f'{lab} تربية'); wt=clone(src_t,f'{lab} تلاوة')
    for i,n in enumerate(names):
        wi.cell(5+i,3).value=n
        wt.cell(4+i,3).value=n
wb.remove(src_i); wb.remove(src_t)
# guide sheet first
gs=wb.create_sheet('طريقة التعبئة',0); gs.sheet_view.rightToLeft=True
rows=['ملف درجات قسم التربية الإسلامية – مدرسة ابن رشد الإعدادية للبنين',
'',
'1) كل شعبة لها ورقتان: "تربية" و"تلاوة"، والأسماء مكتوبة جاهزة.',
'2) عبّ الدرجات في الخانات الفاضية فقط. المجاميع تنحسب تلقائيًا.',
'3) لا تغيّر أسماء الأوراق ولا ترتيب الطلاب.',
'4) خانة "الرقم الشخصي" اتركها فاضية: ما تنرفع للموقع أبدًا.',
'5) بعد التعبئة احفظ الملف وأرسله لكلود، ويحدّث الدرجات في الموقع.',
'',
'تربية: الملاحظة المنظمة (1-4)، ملف الأعمال (1-4)، المهمة (1-4)، الاختبارات (0-20)، تلاوة وحفظ (0-20).',
'تلاوة: الحفظ 1 و2 (0-10)، التلاوة رصد 1 ورصد 2 (0-10 لكل خانة).']
for i,r in enumerate(rows,1):
    c=gs.cell(i,1,r); c.font=Font(name='Arial',size=14 if i==1 else 12,bold=(i==1)); c.alignment=Alignment(horizontal='right')
gs.column_dimensions['A'].width=110
out='/home/claude/grades/درجات_قسم_التربية_الإسلامية.xlsx'
wb.save(out); print('saved',len(wb.sheetnames),wb.sheetnames[:5])
# codes
ALPH='ABCDEFGHJKLMNPQRSTUVWXYZ23456789'
codes={}
try: codes=json.load(open('/home/claude/grades/codes.json',encoding='utf-8'))
except FileNotFoundError: pass
for key,names in ST.items():
    for i,n in enumerate(names):
        k=f'{key}|{i}|{n}'
        if k not in codes:
            c=''.join(secrets.choice(ALPH) for _ in range(8)); codes[k]=c[:4]+'-'+c[4:]
json.dump(codes,open('/home/claude/grades/codes.json','w',encoding='utf-8'),ensure_ascii=False,indent=0)
cw=openpyxl.Workbook(); cs=cw.active; cs.title='رموز الدخول'; cs.sheet_view.rightToLeft=True
cs.append(['الصف','م','اسم الطالب','رمز الدخول'])
for key in sorted(ST):
    g,s=key.split('-')
    for i,n in enumerate(ST[key]):
        cs.append([f'{GN[g]} الإعدادي / {s}',i+1,n,codes[f'{key}|{i}|{n}']])
for col,w in zip('ABCD',(22,6,40,16)): cs.column_dimensions[col].width=w
for c in cs[1]: c.font=Font(name='Arial',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='0B7A52')
for row in cs.iter_rows(min_row=2):
    for c in row: c.font=Font(name='Arial',size=12)
    row[3].font=Font(name='Consolas',size=13,bold=True)
cs.freeze_panes='A2'
cw.save('/home/claude/grades/رموز_دخول_الطلاب.xlsx'); print('codes',len(codes))
